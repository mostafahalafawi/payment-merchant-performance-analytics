"""Automated reporting outputs (openpyxl).

    reports/management_pack.xlsx     formatted multi-sheet pack for leadership
    reports/am_action_list.xlsx      one sheet per account manager: merchants that need action
                                     (declining, dormant / churned, never transacted, negative GP,
                                     billing leakage, low success rate) + a summary sheet
    insights/kpi_summary.md          auto-generated fact sheet

Run:  python python/reporting.py
"""
from __future__ import annotations

import json

import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from utils import get_logger, load_config, read_csv_checked

log = get_logger("reporting")
HEADER_FILL, HEADER_FONT = PatternFill("solid", fgColor="163A5F"), Font(bold=True, color="FFFFFF")
TITLE = Font(bold=True, size=14, color="163A5F")
USD, PCT, PCT2 = '#,##0;[Red]-#,##0', '0.0%', '0.00%'

ACTIONS = {
    "flag_declining": "GMV down >=25% (3M vs prior 3M): call merchant, check competitor pricing / integration issues",
    "flag_inactive": "Dormant / churned: win-back call, confirm integration is still live",
    "flag_never_transacted": "Live but never transacted: activation outreach, test transaction support",
    "flag_negative_gp": "Negative GP: review contract pricing vs processing cost / routing",
    "flag_billing_leakage": "Billed below contract: raise billing correction with Finance, recover arrears",
    "flag_low_success_rate": "Success rate <85% (3M): open ticket with payments ops, review routing / 3DS",
}


def _format(ws, df, fmts, header_row=3):
    for c in range(1, len(df.columns) + 1):
        cell = ws.cell(row=header_row, column=c)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[header_row].height = 32
    for i, col in enumerate(df.columns, start=1):
        width = max([len(str(col))] + [len(str(v)) for v in df[col].head(200)]) if len(df) else len(str(col))
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 48)
        if col in fmts:
            for r in range(header_row + 1, header_row + len(df) + 1):
                ws.cell(row=r, column=i).number_format = fmts[col]
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    if len(df):
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(df.columns))}{header_row + len(df)}"


def _write_book(path, sheets: dict, title_suffix: str, scale_cols=("success_rate", "gp_margin")):
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, (df, _) in sheets.items():
            df.to_excel(xw, sheet_name=name[:31], index=False, startrow=2)
    wb = load_workbook(path)
    for name, (df, fmts) in sheets.items():
        ws = wb[name[:31]]
        ws["A1"] = f"{name} | {title_suffix}"
        ws["A1"].font = TITLE
        _format(ws, df, fmts)
        for col in scale_cols:
            if col in df.columns and len(df):
                L = get_column_letter(list(df.columns).index(col) + 1)
                ws.conditional_formatting.add(f"{L}4:{L}{3 + len(df)}", ColorScaleRule(
                    start_type="percentile", start_value=5, start_color="F8696B", mid_type="percentile", mid_value=50,
                    mid_color="FFEB84", end_type="percentile", end_value=95, end_color="63BE7B"))
    wb.save(path)


def management_pack(cfg, res):
    p, rs = cfg["paths"]["processed"], cfg["paths"]["reports"] / "sql_results"
    cm = read_csv_checked(p / "kpi_company_month.csv")
    dm = read_csv_checked(p / "kpi_dimension_month.csv")
    cty = dm[dm["dimension_name"] == "country"].groupby(["dimension_value", "fiscal_year"], as_index=False)[
        ["gmv", "revenue", "gp", "success_txn", "failed_txn"]].sum()
    cty["gp_margin"] = cty["gp"] / cty["revenue"]
    cty["success_rate"] = cty["success_txn"] / (cty["success_txn"] + cty["failed_txn"])
    trend = cm[["month_key", "gmv", "revenue", "gp", "gp_margin", "take_rate", "success_rate", "refund_rate",
                "active_merchants", "new_live_merchants", "gmv_yoy"]]
    sheets = {
        "Executive KPIs": (read_csv_checked(rs / "q01_executive_kpis_by_fy.csv"), {}),
        "Monthly Trend": (trend, {"gmv": USD, "revenue": USD, "gp": USD, "gp_margin": PCT, "take_rate": PCT2,
                                  "success_rate": PCT2, "refund_rate": PCT2, "gmv_yoy": PCT}),
        "Country": (cty.rename(columns={"dimension_value": "country"}), {"gmv": USD, "revenue": USD, "gp": USD,
                                                                          "gp_margin": PCT, "success_rate": PCT2}),
        "Segment": (read_csv_checked(rs / "q03_segment_growth.csv"), {}),
        "Top Merchants": (read_csv_checked(rs / "q05_top_merchants_gmv.csv"), {}),
        "Profitability": (read_csv_checked(rs / "q06_merchant_profitability.csv"), {}),
        "Growth Bridge": (read_csv_checked(rs / "q17_gmv_growth_bridge.csv"), {}),
        "Gateway Incidents": (read_csv_checked(cfg["paths"]["reports"] / "analysis" / "gateway_incidents.csv"), {}),
        "Billing Leakage": (read_csv_checked(rs / "q16_billing_leakage.csv"), {"leakage_usd": USD}),
        "Data Quality": (read_csv_checked(p / "dq_rule_log.csv"), {}),
    }
    _write_book(cfg["paths"]["reports"] / "management_pack.xlsx", sheets, cfg["project"]["company_name"] + " | synthetic data")
    log.info("management pack -> management_pack.xlsx (%d sheets)", len(sheets))


def am_action_list(cfg):
    h = read_csv_checked(cfg["paths"]["reports"] / "analysis" / "merchant_health.csv")
    flags = list(ACTIONS)
    act = h[h[flags].sum(axis=1) > 0].copy()
    act["recommended_actions"] = act.apply(lambda r: " | ".join(ACTIONS[f] for f in flags if r[f] == 1), axis=1)
    act["priority_gmv_at_stake_usd"] = act[["gmv_prior_3m", "gmv_last_3m"]].max(axis=1) * 4   # annualised run-rate
    cols = ["mid", "merchant_name", "country_code", "segment", "vertical", "health_status", "gmv_prior_3m", "gmv_last_3m",
            "gmv_change_3m_pct", "success_rate_last_3m", "gp_fy26", "leakage_fy26", "priority_gmv_at_stake_usd",
            "recommended_actions"]
    fmts = {"gmv_prior_3m": USD, "gmv_last_3m": USD, "gmv_change_3m_pct": PCT, "success_rate_last_3m": PCT,
            "gp_fy26": USD, "leakage_fy26": USD, "priority_gmv_at_stake_usd": USD}
    summary = act.groupby("account_manager").agg(merchants_needing_action=("mid", "count"),
                                                 **{f: (f, "sum") for f in flags},
                                                 gmv_at_stake_usd=("priority_gmv_at_stake_usd", "sum")).reset_index()
    summary = summary.sort_values("gmv_at_stake_usd", ascending=False)
    sheets = {"Summary": (summary, {"gmv_at_stake_usd": USD})}
    for am, g in act.sort_values("priority_gmv_at_stake_usd", ascending=False).groupby("account_manager", sort=False):
        sheets[am] = (g[cols], fmts)
    _write_book(cfg["paths"]["reports"] / "am_action_list.xlsx", sheets, "Account-manager action list, as of "
                + str(int(h["as_of_month"].iloc[0])), scale_cols=("success_rate_last_3m",))
    log.info("AM action list -> %d merchants across %d account managers", len(act), act["account_manager"].nunique())
    return len(act)


def kpi_summary(cfg, res):
    rs = cfg["paths"]["reports"] / "sql_results"
    q02 = read_csv_checked(rs / "q02_country_contribution.csv")
    q03 = read_csv_checked(rs / "q03_segment_growth.csv")
    L = ["# KPI Summary (auto-generated)", "",
         "_Generated by `python/reporting.py`. Interpretation lives in `business_insights.md`._", "",
         "| KPI | FY2025 | FY2026 | Growth |", "|---|---:|---:|---:|",
         f"| GMV | ${res['gmv_fy25']:,.0f} | ${res['gmv_fy26']:,.0f} | {res['gmv_growth_pct']:+}% |",
         f"| Revenue | ${res['revenue_fy25']:,.0f} | ${res['revenue_fy26']:,.0f} | {res['revenue_growth_pct']:+}% |",
         f"| Gross profit | ${res['gp_fy25']:,.0f} | ${res['gp_fy26']:,.0f} | {res['gp_growth_pct']:+}% |",
         f"| Success rate | {res['success_rate_fy25_pct']}% | {res['success_rate_fy26_pct']}% | |",
         "", f"FY2026 GP margin **{res['gp_margin_fy26_pct']}%**, take rate **{res['take_rate_fy26_pct']}%**.", "",
         "## Countries", "", "| Country | GMV FY26 ($M) | Share | Growth USD | Growth local ccy | GP margin |",
         "|---|---:|---:|---:|---:|---:|"]
    for x in q02.itertuples():
        L.append(f"| {x.country_name} | {x.gmv_fy26_musd} | {x.gmv_share_fy26_pct}% | {x.gmv_growth_usd_pct}% | "
                 f"{x.gmv_growth_local_ccy_pct}% | {x.gp_margin_fy26_pct}% |")
    L += ["", "## Segments", "", "| Segment | GMV FY26 ($M) | Growth | Take rate | GP margin | GP share |", "|---|---:|---:|---:|---:|---:|"]
    for x in q03.itertuples():
        L.append(f"| {x.segment} | {x.gmv_fy26_musd} | {x.gmv_growth_pct}% | {x.take_rate_fy26_pct}% | "
                 f"{x.gp_margin_fy26_pct}% | {x.gp_share_fy26_pct}% |")
    (cfg["paths"]["insights"] / "kpi_summary.md").write_text("\n".join(L) + "\n", encoding="utf-8")


def main():
    cfg = load_config()
    res = json.loads((cfg["paths"]["reports"] / "analysis_results.json").read_text())
    management_pack(cfg, res)
    n = am_action_list(cfg)
    kpi_summary(cfg, res)
    res["am_action_merchants"] = n
    (cfg["paths"]["reports"] / "analysis_results.json").write_text(json.dumps(res, indent=2))


if __name__ == "__main__":
    main()
